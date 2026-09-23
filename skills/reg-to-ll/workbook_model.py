"""Shared, inert data model and schema-aware loader for register workbooks.

The loader deliberately preserves expressions as raw cell values.  Evaluating
those expressions and validating address/bit layouts belong to later stages.
"""

from __future__ import annotations

import ast
import itertools
import operator
import re
from dataclasses import dataclass, replace
from datetime import date, datetime
from pathlib import Path
from typing import Optional, Union

import openpyxl


PathLike = Union[str, Path]
Scalar = Union[str, int, float, bool, date, datetime]

__all__ = [
    "AccessAttribute",
    "FieldModel",
    "HistoryEntry",
    "MetadataEntry",
    "ParameterDefinition",
    "PeripheralModel",
    "ProjectMetadata",
    "RegisterModel",
    "SafeExpressionError",
    "SimplifiedWorkbookError",
    "SourceLocation",
    "WorkbookModel",
    "WorkbookSchemaError",
    "WorkbookValidationError",
    "adapt_simplified_workbook",
    "evaluate_integer_expression",
    "load_workbook_model",
    "validate_and_expand_workbook",
]


class WorkbookSchemaError(ValueError):
    """The workbook schema is mixed, ambiguous, malformed, or unsupported."""


class SimplifiedWorkbookError(WorkbookSchemaError):
    """The workbook cannot be identified as the supported simplified schema."""


class SafeExpressionError(ValueError):
    """地址或复位值表达式不符合受限整数语法。"""


class WorkbookValidationError(ValueError):
    """工作簿的寄存器模型无效。"""


_BINARY_INTEGER_OPERATORS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.FloorDiv: operator.floordiv,
    ast.Mod: operator.mod,
    ast.LShift: operator.lshift,
    ast.RShift: operator.rshift,
    ast.BitOr: operator.or_,
    ast.BitXor: operator.xor,
    ast.BitAnd: operator.and_,
}
_UNARY_INTEGER_OPERATORS = {
    ast.UAdd: operator.pos,
    ast.USub: operator.neg,
    ast.Invert: operator.invert,
}
_MAX_EXPRESSION_LENGTH = 512
_MAX_INTEGER_BITS = 4096


def evaluate_integer_expression(expression, parameters=None):
    """Evaluate a deliberately small integer-only expression language.

    Decimal, hexadecimal, binary and octal literals are supported together
    with parentheses and the operators in the maps above.  A name is usable
    only when the caller explicitly supplies an integer value for it.  Python
    calls, attributes, containers, comparisons and boolean/float values never
    enter the language.
    """
    parameters = {} if parameters is None else dict(parameters)
    if isinstance(expression, bool) or not isinstance(expression, (str, int)):
        raise SafeExpressionError("表达式必须是整数或文本。")
    if isinstance(expression, int):
        return expression
    if not expression.strip():
        raise SafeExpressionError("表达式不能为空。")
    if len(expression) > _MAX_EXPRESSION_LENGTH:
        raise SafeExpressionError("表达式过长。")
    bad_parameter = next(
        (
            name
            for name, value in parameters.items()
            if not isinstance(name, str)
            or not name.isidentifier()
            or isinstance(value, bool)
            or not isinstance(value, int)
        ),
        None,
    )
    if bad_parameter is not None:
        raise SafeExpressionError(f"参数“{bad_parameter}”的名称或整数值无效。")
    try:
        tree = ast.parse(expression, mode="eval")
    except (SyntaxError, ValueError) as error:
        raise SafeExpressionError(f"表达式语法错误：“{expression}”。") from error

    def checked(value):
        if isinstance(value, bool) or not isinstance(value, int):
            raise SafeExpressionError("表达式只能产生整数。")
        if value.bit_length() > _MAX_INTEGER_BITS:
            raise SafeExpressionError("表达式的整数结果过大。")
        return value

    def visit(node, depth=0):
        if depth > 64:
            raise SafeExpressionError("表达式嵌套过深。")
        if isinstance(node, ast.Expression):
            return visit(node.body, depth + 1)
        if isinstance(node, ast.Constant):
            return checked(node.value)
        if isinstance(node, ast.Name):
            if node.id not in parameters:
                raise SafeExpressionError(f"未声明的参数“{node.id}”。")
            return checked(parameters[node.id])
        if isinstance(node, ast.UnaryOp) and type(node.op) in _UNARY_INTEGER_OPERATORS:
            return checked(_UNARY_INTEGER_OPERATORS[type(node.op)](visit(node.operand, depth + 1)))
        if isinstance(node, ast.BinOp) and type(node.op) in _BINARY_INTEGER_OPERATORS:
            left = visit(node.left, depth + 1)
            right = visit(node.right, depth + 1)
            if isinstance(node.op, (ast.LShift, ast.RShift)) and (right < 0 or right > _MAX_INTEGER_BITS):
                raise SafeExpressionError("移位量必须在 0 到 4096 之间。")
            try:
                return checked(_BINARY_INTEGER_OPERATORS[type(node.op)](left, right))
            except (ZeroDivisionError, ValueError) as error:
                raise SafeExpressionError("表达式包含无效的除数或操作数。") from error
        raise SafeExpressionError(
            f"表达式包含不允许的语法“{type(node).__name__}”。"
        )

    return visit(tree)


@dataclass(frozen=True)
class SourceLocation:
    workbook: str
    sheet: Optional[str] = None
    row: Optional[int] = None
    column: Optional[int] = None


@dataclass(frozen=True)
class MetadataEntry:
    key: str
    value: str
    source: SourceLocation
    value_raw: Optional[Scalar] = None


@dataclass(frozen=True)
class HistoryEntry:
    version: str
    record: str
    author: str
    date: str
    date_raw: Optional[Scalar]
    source: SourceLocation


@dataclass(frozen=True)
class ProjectMetadata:
    name: str = ""
    vendor: str = ""
    version: str = ""
    description: str = ""
    entries: tuple[MetadataEntry, ...] = ()
    module_name: str = ""
    bus_type: str = ""
    read_data_reset_port: str = ""
    read_data_reset_value: Optional[int] = None
    read_data_reset_value_raw: Optional[Scalar] = None
    data_width: Optional[int] = None
    data_width_raw: Optional[Scalar] = None
    address_width: Optional[int] = None
    address_width_raw: Optional[Scalar] = None
    author: str = ""
    platform: str = ""
    history: tuple[HistoryEntry, ...] = ()


@dataclass(frozen=True)
class ParameterDefinition:
    name: str
    maximum: Optional[int]
    maximum_raw: Optional[Scalar]
    minimum: Optional[int]
    minimum_raw: Optional[Scalar]
    description: str
    source: SourceLocation


@dataclass(frozen=True)
class AccessAttribute:
    code: str
    cmsis_access: str
    read_action: Optional[str] = None
    modified_write_values: Optional[str] = None


@dataclass(frozen=True)
class FieldModel:
    name: str
    bit_range: str
    bit_range_raw: Optional[Scalar]
    access: AccessAttribute
    reset_value: str
    description: str
    notes: str
    source: SourceLocation
    definition: str = ""
    reset_value_raw: Optional[Scalar] = None
    reset_or_not_raw: Optional[Scalar] = None
    input_or_not_raw: Optional[Scalar] = None
    output_or_not_raw: Optional[Scalar] = None
    hidden_or_not_raw: Optional[Scalar] = None
    boot_address_offset: Optional[int] = None
    boot_address_offset_raw: Optional[Scalar] = None
    write_protection_raw: Optional[Scalar] = None
    read_protection_raw: Optional[Scalar] = None


@dataclass(frozen=True)
class RegisterModel:
    name: str
    address_offset: Optional[int]
    address_offset_raw: Optional[Scalar]
    reset_value: str
    description: str
    notes: str
    fields: tuple[FieldModel, ...]
    source: SourceLocation
    definition: str = ""
    access: Optional[AccessAttribute] = None
    reset_value_raw: Optional[Scalar] = None
    reset_or_not_raw: Optional[Scalar] = None
    input_or_not_raw: Optional[Scalar] = None
    output_or_not_raw: Optional[Scalar] = None
    hidden_or_not_raw: Optional[Scalar] = None
    boot_address_offset: Optional[int] = None
    boot_address_offset_raw: Optional[Scalar] = None
    write_protection_raw: Optional[Scalar] = None
    read_protection_raw: Optional[Scalar] = None


@dataclass(frozen=True)
class PeripheralModel:
    name: str
    registers: tuple[RegisterModel, ...]
    source: SourceLocation


@dataclass(frozen=True)
class WorkbookModel:
    schema: str
    project: ProjectMetadata
    peripherals: tuple[PeripheralModel, ...]
    source: SourceLocation
    parameters: tuple[ParameterDefinition, ...] = ()

    def validate_and_expand(self) -> "WorkbookModel":
        """Validate the complete register layout and expand parameter arrays."""
        return validate_and_expand_workbook(self)


_BIT_RANGE_PATTERN = re.compile(r"^(\d+)(?::(\d+))?$")


def _context(source, register=None, field=None):
    parts = [f"工作簿“{source.workbook}”"]
    if source.sheet:
        parts.append(f"工作表“{source.sheet}”")
    if source.row is not None:
        parts.append(f"第 {source.row} 行")
    if register:
        parts.append(f"寄存器“{register}”")
    if field:
        parts.append(f"字段“{field}”")
    return "、".join(parts)


def _validation_error(message, source, register=None, field=None):
    return WorkbookValidationError(f"{_context(source, register, field)}：{message}")


def _parameter_bound(parameter, value, raw, label):
    if value is not None:
        return value
    try:
        return evaluate_integer_expression(raw)
    except SafeExpressionError as error:
        raise _validation_error(
            f"parameter“{parameter.name}”的维度{label}不是有效整数：{error}",
            parameter.source,
        ) from error


def _expression_names(expression):
    if isinstance(expression, int) and not isinstance(expression, bool):
        return set()
    try:
        tree = ast.parse(str(expression), mode="eval")
    except (SyntaxError, ValueError):
        return set()
    return {node.id for node in ast.walk(tree) if isinstance(node, ast.Name)}


def validate_and_expand_workbook(model: WorkbookModel) -> WorkbookModel:
    """Return a validated model whose register arrays are concrete registers.

    Array dimensions come from parameters referenced by an address expression;
    each inclusive minimum/maximum range is expanded in declaration order.
    """
    data_width = model.project.data_width
    address_width = model.project.address_width
    if not isinstance(data_width, int) or isinstance(data_width, bool) or data_width <= 0 or data_width % 8:
        raise _validation_error("data width 必须是大于 0 且可被 8 整除的整数。", model.source)
    if not isinstance(address_width, int) or isinstance(address_width, bool) or address_width <= 0:
        raise _validation_error("address width 必须是大于 0 的整数。", model.source)
    alignment = data_width // 8

    definitions = {}
    bounds = {}
    for definition in model.parameters:
        key = definition.name
        if not key.isidentifier():
            raise _validation_error(f"parameter 名称“{key}”不是合法标识符。", definition.source)
        if key in definitions:
            raise _validation_error(f"parameter 名称“{key}”重复。", definition.source)
        minimum = _parameter_bound(definition, definition.minimum, definition.minimum_raw, "最小值")
        maximum = _parameter_bound(definition, definition.maximum, definition.maximum_raw, "最大值")
        if minimum > maximum:
            raise _validation_error(
                f"parameter“{key}”的维度最小值 {minimum} 大于最大值 {maximum}。",
                definition.source,
            )
        definitions[key] = definition
        bounds[key] = (minimum, maximum)

    expanded_peripherals = []
    for peripheral in model.peripherals:
        original_names = set()
        expanded_names = set()
        addresses = set()
        previous_address = -1
        expanded_registers = []
        for register in peripheral.registers:
            normalized_name = register.name.casefold()
            if normalized_name in original_names:
                raise _validation_error("寄存器名称重名。", register.source, register.name)
            original_names.add(normalized_name)

            field_names = set()
            occupied_bits = set()
            reset_value = 0
            for item in register.fields:
                normalized_field = item.name.casefold()
                if normalized_field in field_names:
                    raise _validation_error("字段名称重名。", item.source, register.name, item.name)
                field_names.add(normalized_field)
                match = _BIT_RANGE_PATTERN.fullmatch(item.bit_range)
                if match is None:
                    raise _validation_error(
                        "字段位范围格式错误，应为“高位:低位”或单个位号，且不能含空格。",
                        item.source, register.name, item.name,
                    )
                high = int(match.group(1))
                low = int(match.group(2)) if match.group(2) is not None else high
                if high < low:
                    raise _validation_error("字段位范围的高位小于低位。", item.source, register.name, item.name)
                if high >= data_width:
                    raise _validation_error("字段位范围超出寄存器位宽。", item.source, register.name, item.name)
                bits = set(range(low, high + 1))
                if occupied_bits & bits:
                    raise _validation_error("字段位范围与同一寄存器中的其他字段重叠。", item.source, register.name, item.name)
                occupied_bits |= bits
                reset_expression = item.reset_value_raw
                if reset_expression in (None, ""):
                    field_reset = 0
                else:
                    try:
                        field_reset = evaluate_integer_expression(reset_expression)
                    except SafeExpressionError as error:
                        raise _validation_error(f"字段复位值无效：{error}", item.source, register.name, item.name) from error
                width = high - low + 1
                if field_reset < 0 or field_reset >= (1 << width):
                    raise _validation_error("字段复位值超出字段位宽。", item.source, register.name, item.name)
                if not item.name.upper().startswith(("RSV", "RESERVED")):
                    reset_value |= field_reset << low
            if reset_value >= (1 << data_width):
                raise _validation_error("组装后的寄存器复位值超出寄存器位宽。", register.source, register.name)

            raw_address = register.address_offset_raw
            if raw_address in (None, ""):
                raw_address = register.address_offset
            used_names = _expression_names(raw_address)
            varying = [name for name in definitions if name in used_names and bounds[name][0] != bounds[name][1]]
            ranges = [range(bounds[name][0], bounds[name][1] + 1) for name in varying]
            combinations = itertools.product(*ranges) if ranges else [()]
            for values in combinations:
                environment = {name: minimum for name, (minimum, _) in bounds.items()}
                environment.update(dict(zip(varying, values)))
                try:
                    address = evaluate_integer_expression(raw_address, environment)
                except SafeExpressionError as error:
                    raise _validation_error(f"地址表达式无效：{error}", register.source, register.name) from error
                expanded_name = register.name + "".join(str(value) for value in values)
                if expanded_name.casefold() in expanded_names:
                    raise _validation_error("展开后的寄存器名称重复。", register.source, expanded_name)
                if address < 0 or address >= (1 << address_width):
                    raise _validation_error("地址超出地址位宽（address width）的表示范围。", register.source, expanded_name)
                if address % alignment:
                    raise _validation_error(f"地址未按 {alignment} 字节对齐。", register.source, expanded_name)
                if address in addresses:
                    raise _validation_error("寄存器地址重复。", register.source, expanded_name)
                if address <= previous_address:
                    raise _validation_error("寄存器地址必须严格递增。", register.source, expanded_name)
                addresses.add(address)
                previous_address = address
                expanded_names.add(expanded_name.casefold())
                expanded_registers.append(replace(
                    register,
                    name=expanded_name,
                    address_offset=address,
                    address_offset_raw=address,
                    reset_value=hex(reset_value),
                    reset_value_raw=reset_value,
                ))
        expanded_peripherals.append(replace(peripheral, registers=tuple(expanded_registers)))
    return replace(model, peripherals=tuple(expanded_peripherals))


_ACCESS_MAP = {
    "RW": ("read-write", None, None),
    "RO": ("read-only", None, None),
    "WO": ("write-only", None, None),
    "RC": ("read-only", "clear", None),
    "WC": ("read-write", None, "clear"),
    "RS": ("read-only", "set", None),
    "WS": ("read-write", None, "set"),
    "W0C": ("read-write", None, "zeroToClear"),
    "W0S": ("read-write", None, "zeroToSet"),
    "W1C": ("read-write", None, "oneToClear"),
    "W1S": ("read-write", None, "oneToSet"),
    "WRC": ("read-write", "clear", None),
    "WRS": ("read-write", "set", None),
    "WSRC": ("read-write", "clear", "set"),
    "WCRS": ("read-write", "set", "clear"),
    "W1T": ("read-write", None, "oneToToggle"),
    "W0T": ("read-write", None, "zeroToToggle"),
    "W0SRC": ("read-write", "clear", "zeroToSet"),
    "W0CRS": ("read-write", "set", "zeroToClear"),
    "W1SRC": ("read-write", "clear", "oneToSet"),
    "W1CRS": ("read-write", "set", "oneToClear"),
    "WOC": ("write-only", None, "clear"),
    "WOS": ("write-only", None, "set"),
    "W1": ("writeOnce", None, None),
    "WO1": ("writeOnce", None, None),
    "W1W": ("read-write", None, "modify"),
    "WO1W": ("write-only", None, "modify"),
}

_HEADER_ALIASES = (
    {"regname", "registername", "register name"},
    {"addroffset", "addr_offset", "addr offset", "addressoffset", "address offset"},
    {"content", "fieldname", "field name"},
    {"bitrange", "bit_range", "bit range"},
    {"attribute", "access"},
    {"definition"},
    {"defaultvalue", "default_value", "default value", "resetvalue", "reset_value"},
)
_OPTIONAL_EIGHTH_HEADERS = {"description", "notes", "remark", "remarks"}

_REGTOOL_HEADERS = (
    "regname",
    "addroffset",
    "fieldname",
    "bitrange",
    "attribute",
    "definition",
    "resetvalue",
    "description",
    "note",
    "reset_or_not",
    "input_or_not",
    "output_or_not",
    "hiden_nor_not",
    "boot_addr_offset",
    "write_proctect",
    "read_proctect",
)

_INFO_KEY_ALIASES = (
    {"project name", "project_name"},
    {"module name", "module_name"},
    {"bus type", "bus_type"},
    {"rdata reset port", "rdata_reset_port"},
    {"rdata reset value", "rdata_reset_value"},
    {"data width", "data_width"},
    {"addr width", "addr_width"},
    {"author"},
    {"platform"},
)

_PARAMETER_HEADERS = (
    {"parameter_name", "parameter name"},
    {"max_value", "max value"},
    {"min_value", "min value"},
    {"discription", "description"},
)


def _text(value, default=""):
    return default if value is None else str(value).strip()


def _normalized_header(value):
    return " ".join(_text(value).lower().split())


def _normalized_regtool_header(value):
    """Normalize harmless case/whitespace only; do not repair misspellings."""
    return "".join(_text(value).lower().split())


def _schema_for_headers(values):
    headers = [_normalized_header(value) for value in values]
    while headers and not headers[-1]:
        headers.pop()
    if len(headers) not in (7, 8):
        return None
    if any(headers[index] not in aliases for index, aliases in enumerate(_HEADER_ALIASES)):
        return None
    if len(headers) == 8 and headers[7] not in _OPTIONAL_EIGHTH_HEADERS:
        return None
    return f"simplified-{len(headers)}"


def _is_regtool_header(values):
    headers = [_normalized_regtool_header(value) for value in values]
    while headers and not headers[-1]:
        headers.pop()
    return tuple(headers) == _REGTOOL_HEADERS


def _suspected_regtool_header_error(values):
    headers = [_normalized_regtool_header(value) for value in values]
    while headers and not headers[-1]:
        headers.pop()
    if len(headers) <= 8:
        return None
    matches = sum(
        index < len(headers) and headers[index] == expected
        for index, expected in enumerate(_REGTOOL_HEADERS)
    )
    if matches < 6:
        return None
    for index, expected in enumerate(_REGTOOL_HEADERS):
        if index >= len(headers) or headers[index] != expected:
            return index + 1, expected
    if len(headers) > len(_REGTOOL_HEADERS):
        return 17, "表头应在第 16 列结束"
    return None


def _suspected_register_header_error(values):
    """Return the first bad header position for a register-like sheet.

    A sheet is register-like when at least three of the seven required positions
    use known aliases.  This still recognizes a table whose first header is
    misspelled, while keeping ordinary README/calculation sheets ignorable.
    """
    headers = [_normalized_header(value) for value in values]
    while headers and not headers[-1]:
        headers.pop()
    positional_matches = sum(
        index < len(headers) and headers[index] in aliases
        for index, aliases in enumerate(_HEADER_ALIASES)
    )
    if not headers or positional_matches < 3:
        return None
    for index, aliases in enumerate(_HEADER_ALIASES):
        if index >= len(headers) or headers[index] not in aliases:
            return index + 1
    if len(headers) >= 8 and headers[7] not in _OPTIONAL_EIGHTH_HEADERS:
        return 8
    if len(headers) > 8:
        return 9
    return None


def _parse_integer(value):
    text = _text(value)
    if not text or "+" in text or "*" in text:
        return None
    try:
        return int(text, 16) if text.lower().startswith("0x") else int(text)
    except ValueError:
        return None


def _safe_scalar(value):
    if value is None or isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    return str(value)


def _project_metadata(workbook, workbook_path):
    entries = []
    history = []
    for sheet in workbook.worksheets:
        if sheet.title.lower() != "info":
            continue
        in_history = False
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            first = _normalized_header(row[0]) if row else ""
            if first in {"history log", "history_log"}:
                in_history = True
                continue
            if in_history:
                if first == "version":
                    continue
                if not first:
                    continue
                cells = list(row) + [None] * max(0, 4 - len(row))
                history.append(HistoryEntry(
                    version=_text(cells[0]),
                    record=_text(cells[1]),
                    author=_text(cells[2]),
                    date=_text(cells[3]),
                    date_raw=_safe_scalar(cells[3]),
                    source=SourceLocation(workbook_path, sheet.title, row_number, 1),
                ))
                continue
            if len(row) < 2 or row[0] is None or row[1] is None:
                continue
            entries.append(MetadataEntry(
                key=_text(row[0]),
                value=_text(row[1]),
                source=SourceLocation(workbook_path, sheet.title, row_number, 1),
                value_raw=_safe_scalar(row[1]),
            ))
        break
    normalized = {_normalized_header(entry.key): entry for entry in entries}

    def value(*keys):
        entry = next((normalized[key] for key in keys if key in normalized), None)
        return entry.value if entry else ""

    def raw(*keys):
        entry = next((normalized[key] for key in keys if key in normalized), None)
        return entry.value_raw if entry else None

    return ProjectMetadata(
        name=value("project name", "project_name", "project", "name"),
        vendor=value("vendor"),
        version=value("version"),
        description=value("description"),
        entries=tuple(entries),
        module_name=value("module name", "module_name"),
        bus_type=value("bus type", "bus_type"),
        read_data_reset_port=value("rdata reset port", "rdata_reset_port"),
        read_data_reset_value=_parse_integer(raw("rdata reset value", "rdata_reset_value")),
        read_data_reset_value_raw=raw("rdata reset value", "rdata_reset_value"),
        data_width=_parse_integer(raw("data width", "data_width")),
        data_width_raw=raw("data width", "data_width"),
        address_width=_parse_integer(raw("addr width", "addr_width")),
        address_width_raw=raw("addr width", "addr_width"),
        author=value("author"),
        platform=value("platform"),
        history=tuple(history),
    )


def _access_attribute(value):
    code = _text(value, "RW") or "RW"
    cmsis_access, read_action, modified_write_values = _ACCESS_MAP.get(
        code.upper(), ("read-write", None, None)
    )
    return AccessAttribute(code, cmsis_access, read_action, modified_write_values)


def _parse_peripheral(sheet, workbook_path):
    registers = []
    current = None
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        cells = list(row) + [None] * max(0, 8 - len(row))
        register_name = _text(cells[0])
        field_name = _text(cells[2])
        bit_range = _text(cells[3])
        if register_name:
            if current is not None:
                registers.append(RegisterModel(fields=tuple(current.pop("fields")), **current))
            current = {
                "name": register_name,
                "address_offset": _parse_integer(cells[1]),
                "address_offset_raw": _safe_scalar(cells[1]),
                "reset_value": _text(cells[6]),
                "description": _text(cells[5]),
                "notes": _text(cells[7]),
                "fields": [],
                "source": SourceLocation(workbook_path, sheet.title, row_number, 1),
                "definition": _text(cells[5]),
                "access": _access_attribute(cells[4]) if cells[4] is not None else None,
                "reset_value_raw": _safe_scalar(cells[6]),
            }
        if current is not None and field_name and bit_range:
            current["fields"].append(FieldModel(
                name=field_name,
                bit_range=bit_range,
                bit_range_raw=_safe_scalar(cells[3]),
                access=_access_attribute(cells[4]),
                reset_value=_text(cells[6]),
                description=_text(cells[5]),
                notes=_text(cells[7]),
                source=SourceLocation(workbook_path, sheet.title, row_number, 3),
                definition=_text(cells[5]),
                reset_value_raw=_safe_scalar(cells[6]),
            ))
    if current is not None:
        registers.append(RegisterModel(fields=tuple(current.pop("fields")), **current))
    return PeripheralModel(
        name=sheet.title,
        registers=tuple(registers),
        source=SourceLocation(workbook_path, sheet.title, 1, 1),
    )


def _raw_register_columns(cells):
    return {
        "reset_or_not_raw": _safe_scalar(cells[9]),
        "input_or_not_raw": _safe_scalar(cells[10]),
        "output_or_not_raw": _safe_scalar(cells[11]),
        "hidden_or_not_raw": _safe_scalar(cells[12]),
        "boot_address_offset": _parse_integer(cells[13]),
        "boot_address_offset_raw": _safe_scalar(cells[13]),
        "write_protection_raw": _safe_scalar(cells[14]),
        "read_protection_raw": _safe_scalar(cells[15]),
    }


def _parse_regtool_peripheral(sheet, workbook_path):
    registers = []
    current = None
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        cells = list(row) + [None] * max(0, 16 - len(row))
        register_name = _text(cells[0])
        field_name = _text(cells[2])
        bit_range = _text(cells[3])
        if register_name:
            if current is not None:
                registers.append(RegisterModel(fields=tuple(current.pop("fields")), **current))
            current = {
                "name": register_name,
                "address_offset": _parse_integer(cells[1]),
                "address_offset_raw": _safe_scalar(cells[1]),
                "reset_value": _text(cells[6]),
                "description": _text(cells[7]),
                "notes": _text(cells[8]),
                "fields": [],
                "source": SourceLocation(workbook_path, sheet.title, row_number, 1),
                "definition": _text(cells[5]),
                "access": _access_attribute(cells[4]) if cells[4] is not None else None,
                "reset_value_raw": _safe_scalar(cells[6]),
                **_raw_register_columns(cells),
            }
        if current is not None and field_name and bit_range:
            current["fields"].append(FieldModel(
                name=field_name,
                bit_range=bit_range,
                bit_range_raw=_safe_scalar(cells[3]),
                access=_access_attribute(cells[4]),
                reset_value=_text(cells[6]),
                description=_text(cells[7]),
                notes=_text(cells[8]),
                source=SourceLocation(workbook_path, sheet.title, row_number, 3),
                definition=_text(cells[5]),
                reset_value_raw=_safe_scalar(cells[6]),
                **_raw_register_columns(cells),
            ))
    if current is not None:
        registers.append(RegisterModel(fields=tuple(current.pop("fields")), **current))
    return PeripheralModel(
        name=sheet.title,
        registers=tuple(registers),
        source=SourceLocation(workbook_path, sheet.title, 1, 1),
    )


def _validate_regtool_info(sheet):
    for row_number, aliases in enumerate(_INFO_KEY_ALIASES, start=1):
        actual = _normalized_header(sheet.cell(row_number, 1).value)
        if actual not in aliases:
            expected = sorted(aliases)[0]
            raise WorkbookSchemaError(
                f"regtool 的 info 工作表已损坏：第 {row_number} 行第 1 列"
                f"应为“{expected}”，实际为“{_text(sheet.cell(row_number, 1).value)}”。"
            )


def _parse_parameters(sheet, workbook_path):
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    for index, aliases in enumerate(_PARAMETER_HEADERS):
        actual = _normalized_header(header[index] if index < len(header) else None)
        if actual not in aliases:
            expected = sorted(aliases)[0]
            raise WorkbookSchemaError(
                f"regtool 的 parameter 工作表已损坏：第 1 行第 {index + 1} 列"
                f"应为“{expected}”，实际为“{_text(header[index] if index < len(header) else None)}”。"
            )
    for index, value in enumerate(header[4:], start=5):
        if _text(value):
            raise WorkbookSchemaError(
                f"regtool 的 parameter 工作表已损坏：第 1 行第 {index} 列"
                "超出约定的 4 列格式。"
            )
    definitions = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        cells = list(row) + [None] * max(0, 4 - len(row))
        if all(value is None for value in cells[:4]):
            continue
        if not _text(cells[0]):
            raise WorkbookSchemaError(
                f"parameter 工作表第 {row_number} 行第 1 列缺少 parameter_name。"
            )
        definitions.append(ParameterDefinition(
            name=_text(cells[0]),
            maximum=_parse_integer(cells[1]),
            maximum_raw=_safe_scalar(cells[1]),
            minimum=_parse_integer(cells[2]),
            minimum_raw=_safe_scalar(cells[2]),
            description=_text(cells[3]),
            source=SourceLocation(workbook_path, sheet.title, row_number, 1),
        ))
    return tuple(definitions)


def load_workbook_model(excel_path: PathLike) -> WorkbookModel:
    """Load one explicitly identified workbook schema into ``WorkbookModel``.

    Supported inputs are the complete historical 16-column regtool workbook
    and the established simplified 7/8-column sheets.  Similar-looking,
    damaged, mixed, or ambiguously ordered workbooks fail at this boundary.
    """
    workbook_path = str(excel_path)
    workbook = openpyxl.load_workbook(excel_path, data_only=False)
    try:
        info_sheets = [sheet for sheet in workbook.worksheets if sheet.title.lower() == "info"]
        parameter_sheets = [
            sheet for sheet in workbook.worksheets if sheet.title.lower() == "parameter"
        ]
        if len(info_sheets) > 1 or len(parameter_sheets) > 1:
            raise WorkbookSchemaError("info 或 parameter 工作表重名，无法明确识别工作簿结构。")

        classified = []
        malformed = []
        for index, sheet in enumerate(workbook.worksheets):
            if sheet.title.lower() in {"info", "parameter"}:
                continue
            header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            if _is_regtool_header(header):
                classified.append((index, sheet, "regtool-16"))
                continue
            schema = _schema_for_headers(header)
            if schema:
                classified.append((index, sheet, schema))
                continue
            regtool_error = _suspected_regtool_header_error(header)
            if regtool_error is not None:
                bad_column, expected = regtool_error
                malformed.append((sheet, bad_column, expected, "regtool"))
                continue
            bad_column = _suspected_register_header_error(header)
            if bad_column is not None:
                malformed.append((sheet, bad_column, "简化 7/8 列表头", "simplified"))

        full = [(index, sheet) for index, sheet, schema in classified if schema == "regtool-16"]
        simplified = [
            (index, sheet, schema)
            for index, sheet, schema in classified
            if schema.startswith("simplified-")
        ]

        if full:
            if simplified:
                full_names = "、".join(sheet.title for _, sheet in full)
                simple_names = "、".join(sheet.title for _, sheet, _ in simplified)
                raise WorkbookSchemaError(
                    f"同一工作簿不能混用 regtool 16 列与简化 schema："
                    f"16 列工作表 {full_names}；简化工作表 {simple_names}。"
                )
            if len(info_sheets) != 1 or len(parameter_sheets) != 1:
                raise WorkbookSchemaError(
                    "regtool 16 列工作簿必须各有一个 info 和 parameter 工作表。"
                )
            info_index = workbook.worksheets.index(info_sheets[0])
            parameter_index = workbook.worksheets.index(parameter_sheets[0])
            if info_index >= parameter_index:
                raise WorkbookSchemaError("regtool 工作簿中 info 必须位于 parameter 之前。")
            outside = [
                sheet.title for index, sheet in full
                if not info_index < index < parameter_index
            ]
            if outside:
                raise WorkbookSchemaError(
                    f"regtool 模块工作表 {('、'.join(outside))} 必须位于 info 与 parameter 之间。"
                )
            full_by_index = {index: sheet for index, sheet in full}
            for index in range(info_index + 1, parameter_index):
                if index not in full_by_index:
                    sheet = workbook.worksheets[index]
                    detail = next(
                        (
                            f"第 1 行第 {column} 列应为“{expected}”"
                            for malformed_sheet, column, expected, _ in malformed
                            if malformed_sheet is sheet
                        ),
                        "未使用完整 16 列表头",
                    )
                    raise WorkbookSchemaError(
                        f"info 与 parameter 之间的工作表“{sheet.title}”格式损坏或存在歧义：{detail}。"
                    )
            if malformed:
                sheet, column, expected, _ = malformed[0]
                raise WorkbookSchemaError(
                    f"疑似寄存器工作表“{sheet.title}”的表头已损坏或存在歧义："
                    f"第 1 行第 {column} 列应为“{expected}”。"
                )
            _validate_regtool_info(info_sheets[0])
            parameters = _parse_parameters(parameter_sheets[0], workbook_path)
            module_sheets = [full_by_index[index] for index in sorted(full_by_index)]
            return WorkbookModel(
                schema="regtool-16",
                project=_project_metadata(workbook, workbook_path),
                peripherals=tuple(
                    _parse_regtool_peripheral(sheet, workbook_path) for sheet in module_sheets
                ),
                source=SourceLocation(workbook_path),
                parameters=parameters,
            )

        if malformed:
            sheet, column, expected, kind = malformed[0]
            if kind == "regtool":
                raise WorkbookSchemaError(
                    f"疑似 regtool 寄存器工作表“{sheet.title}”的表头已损坏或存在歧义："
                    f"第 1 行第 {column} 列应为“{expected}”。"
                )
            raise WorkbookSchemaError(
                f"疑似寄存器工作表“{sheet.title}”的表头已损坏或存在歧义："
                f"第 1 行第 {column} 列不符合约定的 7 列或 8 列格式。"
            )
        if not simplified:
            raise WorkbookSchemaError(
                "无法识别简化 Excel：请使用约定顺序的 7 列或 8 列表头"
                "（register name、addr_offset、content、bit range、attribute、"
                "definition、default_value，以及可选 description）。"
            )
        schemas = {schema for _, _, schema in simplified}
        if len(schemas) != 1:
            raise WorkbookSchemaError("同一工作簿中不能混用简化 7 列和 8 列表头。")
        return WorkbookModel(
            schema=next(iter(schemas)),
            project=_project_metadata(workbook, workbook_path),
            peripherals=tuple(
                _parse_peripheral(sheet, workbook_path) for _, sheet, _ in simplified
            ),
            source=SourceLocation(workbook_path),
            parameters=(
                _parse_parameters(parameter_sheets[0], workbook_path)
                if parameter_sheets else ()
            ),
        )
    finally:
        workbook.close()


def adapt_simplified_workbook(excel_path: PathLike) -> WorkbookModel:
    """Compatibility adapter that accepts only simplified 7/8-column inputs."""
    try:
        model = load_workbook_model(excel_path)
    except WorkbookSchemaError as error:
        raise SimplifiedWorkbookError(str(error)) from error
    if not model.schema.startswith("simplified-"):
        raise SimplifiedWorkbookError(
            "该入口仅支持简化 7/8 列工作簿；完整 regtool 请使用 load_workbook_model。"
        )
    return model
