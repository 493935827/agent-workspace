"""Converged, offline input boundary for local and Feishu workbooks."""

from __future__ import annotations

import json
import tempfile
from dataclasses import replace
from pathlib import Path
from typing import Mapping, Union

import openpyxl

from workbook_model import SourceLocation, WorkbookModel, load_workbook_model


PathLike = Union[str, Path]

__all__ = [
    "InputConvergenceError",
    "load_feishu_workbook_model",
    "load_local_workbook_model",
]


class InputConvergenceError(ValueError):
    """飞书原始快照与重建 Excel 在生成前不一致。"""


def _neutral_source(source: SourceLocation) -> SourceLocation:
    return replace(source, workbook="<normalized-workbook>")


def _normalize_sources(model: WorkbookModel) -> WorkbookModel:
    project = replace(
        model.project,
        entries=tuple(replace(item, source=_neutral_source(item.source)) for item in model.project.entries),
        history=tuple(replace(item, source=_neutral_source(item.source)) for item in model.project.history),
    )
    peripherals = tuple(
        replace(
            peripheral,
            source=_neutral_source(peripheral.source),
            registers=tuple(
                replace(
                    register,
                    source=_neutral_source(register.source),
                    fields=tuple(
                        replace(field, source=_neutral_source(field.source))
                        for field in register.fields
                    ),
                )
                for register in peripheral.registers
            ),
        )
        for peripheral in model.peripherals
    )
    parameters = tuple(
        replace(item, source=_neutral_source(item.source)) for item in model.parameters
    )
    return replace(
        model,
        project=project,
        peripherals=peripherals,
        parameters=parameters,
        source=_neutral_source(model.source),
    )


def load_local_workbook_model(excel_path: PathLike) -> WorkbookModel:
    """通过唯一的 schema/model/validation 边界加载本地 Excel。"""
    return _normalize_sources(load_workbook_model(excel_path).validate_and_expand())


def _plain_cell(raw):
    if not isinstance(raw, list):
        return raw
    pieces = []
    for item in raw:
        if isinstance(item, Mapping):
            pieces.append(str(item.get("text", "")))
        else:
            pieces.append(str(item))
    return "".join(pieces) or None


def _snapshot_data(snapshot: Union[PathLike, Mapping]) -> Mapping:
    if isinstance(snapshot, Mapping):
        return snapshot
    try:
        return json.loads(Path(snapshot).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise InputConvergenceError(f"无法读取飞书原始 JSON 快照：{error}") from error


def _snapshot_workbook(snapshot: Union[PathLike, Mapping], output: Path) -> None:
    data = _snapshot_data(snapshot)
    sheets = data.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise InputConvergenceError("飞书原始 JSON 快照中没有工作表。")
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    try:
        for sheet_number, item in enumerate(sheets, start=1):
            if not isinstance(item, Mapping) or not isinstance(item.get("values"), list):
                raise InputConvergenceError(f"飞书原始 JSON 快照第 {sheet_number} 个工作表格式无效。")
            title = str(item.get("title") or f"Sheet{sheet_number}")
            sheet = workbook.create_sheet(title)
            for row in item["values"]:
                if not isinstance(row, list):
                    raise InputConvergenceError(f"飞书工作表“{title}”包含无效行。")
                sheet.append([_plain_cell(cell) for cell in row])
        workbook.save(output)
    finally:
        workbook.close()


def _semantic_difference(raw: WorkbookModel, rebuilt: WorkbookModel) -> str | None:
    raw_peripherals = [item.name for item in raw.peripherals]
    rebuilt_peripherals = [item.name for item in rebuilt.peripherals]
    if raw_peripherals != rebuilt_peripherals:
        return f"模块顺序或集合不一致（飞书={raw_peripherals}，Excel={rebuilt_peripherals}）。"
    for raw_peripheral, excel_peripheral in zip(raw.peripherals, rebuilt.peripherals):
        raw_names = [item.name for item in raw_peripheral.registers]
        excel_names = [item.name for item in excel_peripheral.registers]
        missing = [name for name in raw_names if name not in excel_names]
        extra = [name for name in excel_names if name not in raw_names]
        if missing:
            return f"模块“{raw_peripheral.name}”的重建 Excel 缺少寄存器：{missing}。"
        if extra:
            return f"模块“{raw_peripheral.name}”的重建 Excel 多出寄存器：{extra}。"
        if raw_names != excel_names:
            return f"模块“{raw_peripheral.name}”的寄存器顺序不一致。"
        for raw_register, excel_register in zip(raw_peripheral.registers, excel_peripheral.registers):
            if (raw_register.name, raw_register.address_offset_raw) != (
                excel_register.name, excel_register.address_offset_raw
            ):
                return f"寄存器“{raw_register.name}”的地址语义不一致。"
            raw_fields = [field.name for field in raw_register.fields]
            excel_fields = [field.name for field in excel_register.fields]
            if raw_fields != excel_fields:
                return f"寄存器“{raw_register.name}”的字段顺序或集合不一致。"
            for raw_field, excel_field in zip(raw_register.fields, excel_register.fields):
                raw_semantics = (
                    raw_field.name, raw_field.bit_range, raw_field.access,
                    raw_field.reset_value, raw_field.definition,
                )
                excel_semantics = (
                    excel_field.name, excel_field.bit_range, excel_field.access,
                    excel_field.reset_value, excel_field.definition,
                )
                if raw_semantics != excel_semantics:
                    return f"寄存器“{raw_register.name}”字段“{raw_field.name}”语义不一致。"
    return None


def load_feishu_workbook_model(
    raw_snapshot: Union[PathLike, Mapping], rebuilt_excel_path: PathLike
) -> WorkbookModel:
    """核对飞书快照与重建 Excel，然后返回与本地入口相同的规范模型。"""
    with tempfile.TemporaryDirectory(prefix="regtool-feishu-") as directory:
        raw_excel = Path(directory) / "raw-snapshot.xlsx"
        _snapshot_workbook(raw_snapshot, raw_excel)
        raw_model = load_workbook_model(raw_excel)
    rebuilt_model = load_workbook_model(rebuilt_excel_path)
    difference = _semantic_difference(raw_model, rebuilt_model)
    if difference:
        raise InputConvergenceError(f"飞书快照校验失败：{difference} 已在生成前停止。")
    raw_normalized = _normalize_sources(raw_model.validate_and_expand())
    rebuilt_normalized = _normalize_sources(rebuilt_model.validate_and_expand())
    if raw_normalized != rebuilt_normalized:
        raise InputConvergenceError("飞书快照校验失败：完整规范模型不一致，已在生成前停止。")
    return rebuilt_normalized
